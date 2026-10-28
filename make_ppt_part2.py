
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DEEP_PURPLE   = RGBColor(0x4A, 0x14, 0x8C)
MID_PURPLE    = RGBColor(0x7B, 0x1F, 0xA2)
ACCENT_TEAL   = RGBColor(0x00, 0x89, 0x7B)
ACCENT_RED    = RGBColor(0xC6, 0x28, 0x28)
ACCENT_ORANGE = RGBColor(0xEF, 0x6C, 0x00)
ACCENT_GREEN  = RGBColor(0x2E, 0x7D, 0x32)
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT     = RGBColor(0x21, 0x21, 0x21)

prs = Presentation("temp_part1.pptx")
BLANK = prs.slide_layouts[6]

def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def rect(slide, l, t, w, h, fill_color):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape

def txt(slide, text, l, t, w, h, size=18, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p  = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return tb

def bullet_box(slide, items, l, t, w, h, size=15, color=DARK_TEXT, bullet="▸ "):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = bullet + item
        run.font.size = Pt(size)
        run.font.color.rgb = color

def divider(slide, t, color=MID_PURPLE, thickness=3):
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(t), Inches(12.33), Inches(thickness/72))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — System Architecture (pipeline diagram)
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s, RGBColor(0xF3, 0xE5, 0xF5))
rect(s, 0, 0, 13.33, 1.3, DEEP_PURPLE)
txt(s, "System Architecture", 0.5, 0.2, 12.0, 0.9,
    size=34, bold=True, color=WHITE)

# Pipeline boxes — horizontal flow
steps = [
    ("📨\nMessage /\nScreenshot",   ACCENT_TEAL),
    ("🔤\nLanguage\nDetection",     MID_PURPLE),
    ("🤖\nAI / NLP\nAnalysis",      DEEP_PURPLE),
    ("🔒\nSecurity\nRules",         ACCENT_ORANGE),
    ("🔗\nURL\nIntelligence",       ACCENT_RED),
    ("⚖️\nRisk\nEngine",            ACCENT_GREEN),
]
box_w, box_h = 1.75, 1.55
start_x = 0.35
for i, (label, col) in enumerate(steps):
    cx = start_x + i * 2.15
    rect(s, cx, 1.55, box_w, box_h, col)
    txt(s, label, cx+0.05, 1.6, box_w-0.1, box_h-0.1,
        size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    if i < len(steps)-1:
        txt(s, "→", cx+box_w, 2.05, 0.35, 0.6,
            size=22, bold=True, color=DEEP_PURPLE, align=PP_ALIGN.CENTER)

# Output card
rect(s, 1.5, 3.4, 10.33, 2.5, RGBColor(0x1A, 0x23, 0x7E))
txt(s, "Risk Engine Output", 1.7, 3.45, 9.9, 0.5,
    size=16, bold=True, color=RGBColor(0xFF,0xF1,0x76), align=PP_ALIGN.CENTER)

output_cols = [
    ("Risk Score\n0 – 100", "Threat Type\nTELEBIRR_SCAM\nBANK_PHISHING…"),
    ("Confidence\n0.0 – 1.0", "Language\nENGLISH / AMHARIC\n/ MIXED"),
    ("Indicators\nurgency · credential\n· lookalike_domain", "Recommendation\n+ Explanation"),
]
for i, (top, bot) in enumerate(output_cols):
    cx = 1.7 + i * 3.35
    rect(s, cx, 3.95, 3.1, 0.55, RGBColor(0x28, 0x35, 0x93))
    txt(s, top, cx+0.05, 3.97, 3.0, 0.5, size=11, color=WHITE, align=PP_ALIGN.CENTER)
    rect(s, cx, 4.55, 3.1, 0.65, RGBColor(0x1A, 0x23, 0x7E))
    txt(s, bot, cx+0.05, 4.57, 3.0, 0.62, size=10,
        color=RGBColor(0xC5,0xCA,0xFF), align=PP_ALIGN.CENTER)

# Detect→Explain→Educate→Protect banner
rect(s, 0.5, 6.1, 12.33, 0.65, ACCENT_TEAL)
txt(s, "Detect  →  Explain  →  Educate  →  Protect",
    0.5, 6.15, 12.33, 0.55, size=18, bold=True,
    color=WHITE, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — What Is Built vs What Remains
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s, RGBColor(0xFA, 0xF5, 0xFF))
rect(s, 0, 0, 13.33, 1.3, DEEP_PURPLE)
txt(s, "What Is Built  vs  What Remains", 0.5, 0.2, 12.0, 0.9,
    size=32, bold=True, color=WHITE)

# BUILT column
rect(s, 0.4, 1.45, 5.9, 0.55, ACCENT_GREEN)
txt(s, "✅  BUILT & WORKING", 0.4, 1.45, 5.9, 0.55,
    size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
built = [
    "FastAPI backend  (POST /analyze)",
    "Multilingual security rules engine (EN + AM)",
    "URL intelligence: shorteners, lookalike domains, IP URLs",
    "Unicode language & code-switch detector",
    "Composite risk engine (rules + ML score)",
    "10 threat categories with explanations",
    "Trained ML baseline (TF-IDF + Logistic Regression)",
    "Flutter mobile app — EN & Amharic UI",
    "Dataset (raw + preprocessed)",
    "Evaluation & preprocessing pipeline scripts",
]
bullet_box(s, built, 0.4, 2.05, 5.9, 4.8, size=13,
           color=DARK_TEXT, bullet="✔ ")

# REMAINING column
rect(s, 6.9, 1.45, 5.9, 0.55, ACCENT_ORANGE)
txt(s, "🔲  REMAINING / FUTURE", 6.9, 1.45, 5.9, 0.55,
    size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
remaining = [
    "OCR / screenshot input (camera → text)",
    "Share-to-analyze from WhatsApp / Telegram",
    "Larger real-world Amharic dataset",
    "Transformer-based multilingual model",
    "Live threat-intelligence API integration",
    "On-device inference (privacy-first)",
    "Afaan Oromoo language support",
    "Education / awareness module in app",
    "SMS channel listener",
    "Production deployment & auth layer",
]
bullet_box(s, remaining, 6.9, 2.05, 5.9, 4.8, size=13,
           color=DARK_TEXT, bullet="○ ")

divider(s, 7.1, ACCENT_TEAL, 2)
txt(s, "Core detection engine is fully functional — remaining items are channel extensions & scale improvements",
    0.5, 7.15, 12.33, 0.35, size=12, italic=True,
    color=MID_PURPLE, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — Live Demo Walkthrough
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s, RGBColor(0x1A, 0x00, 0x38))
rect(s, 0, 0, 13.33, 1.3, MID_PURPLE)
txt(s, "Live Demo — How It Works", 0.5, 0.2, 12.0, 0.9,
    size=34, bold=True, color=WHITE)

# Step flow
steps5 = [
    ("1", "User pastes\nsuspicious message\n(EN / AM / Mixed)",   ACCENT_TEAL),
    ("2", "Flutter app\nsends POST /analyze\nto FastAPI",          MID_PURPLE),
    ("3", "Language detected\nRules + URL\nanalysis run",          ACCENT_ORANGE),
    ("4", "ML model scores\nRisk engine\ncombines signals",        ACCENT_RED),
    ("5", "Result returned:\nScore · Type\nIndicators · Action",   ACCENT_GREEN),
]
for i, (num, label, col) in enumerate(steps5):
    cx = 0.5 + i * 2.55
    rect(s, cx, 1.55, 2.2, 2.0, col)
    txt(s, num,   cx+0.05, 1.6,  2.1, 0.5, size=28, bold=True,
        color=WHITE, align=PP_ALIGN.CENTER)
    txt(s, label, cx+0.05, 2.1,  2.1, 1.3, size=12,
        color=WHITE, align=PP_ALIGN.CENTER)
    if i < 4:
        txt(s, "→", cx+2.2, 2.2, 0.3, 0.6, size=20, bold=True,
            color=RGBColor(0xCE,0x93,0xD8), align=PP_ALIGN.CENTER)

# Sample message box
rect(s, 0.5, 3.85, 12.33, 0.45, RGBColor(0x2D, 0x00, 0x5E))
txt(s, "Sample message sent to the API:", 0.6, 3.88, 12.0, 0.38,
    size=13, bold=True, color=RGBColor(0xFF,0xF1,0x76))

rect(s, 0.5, 4.35, 12.33, 0.75, RGBColor(0x0D, 0x0D, 0x0D))
txt(s,
    '"የእርስዎ telebirr ሂሳብ ታግዷል። PIN ቁጥርዎን በ http://cbe-verify.top/login ያስገቡ"',
    0.65, 4.38, 12.0, 0.65, size=13,
    color=RGBColor(0x80, 0xFF, 0x80), italic=True)

# Expected output
rect(s, 0.5, 5.25, 12.33, 0.45, RGBColor(0x2D, 0x00, 0x5E))
txt(s, "Expected API response:", 0.6, 5.28, 12.0, 0.38,
    size=13, bold=True, color=RGBColor(0xFF,0xF1,0x76))

rect(s, 0.5, 5.75, 12.33, 1.35, RGBColor(0x0D, 0x0D, 0x0D))
txt(s,
    'classification: HIGH_RISK  |  risk_score: 90  |  threat_type: TELEBIRR_SCAM\n'
    'language: MIXED  |  code_switched: true\n'
    'indicators: [credential_request, fear_or_threat, lookalike_domain, suspicious_tld, credential_path]\n'
    'recommendation: "Never share your 6-digit telebirr PIN or OTP with anyone."',
    0.65, 5.78, 12.0, 1.28, size=11,
    color=RGBColor(0x80, 0xFF, 0x80))

prs.save("temp_part2.pptx")
print("part2 saved")
